"""Resume text extraction and Markdown to DOCX / PDF export.

The AI assistant writes cover letters and tailored resumes as Markdown in the
workspace. `jobpilot export file.md --format docx|pdf` turns them into files you
can send. Supported Markdown: # headings, **bold**, *italic*, - bullets, blank
lines between paragraphs.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import List, Tuple

TECH_TERMS = ["Python", "Java", "TypeScript", "JavaScript", "Go", "Rust", "C++", "C#", "Kotlin", "Swift", "SQL",
              "FastAPI", "Django", "Flask", "Spring", "Node.js", "NestJS", "Express", "React", "Next.js", "Vue",
              "Angular", "PostgreSQL", "MySQL", "MongoDB", "Redis", "Kafka", "Docker", "Kubernetes", "AWS", "Azure",
              "GCP", "Terraform", "Linux", "Git", "GraphQL", "REST", "LangChain", "LangGraph", "CrewAI", "RAG",
              "MCP", "PyTorch", "TensorFlow", "scikit-learn", "Pandas", "NumPy", "LLM", "Machine Learning",
              "Spark", "Airflow", "dbt", "Tableau", "Power BI", "CI/CD", "GitHub Actions"]


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text.strip() for c in r.cells))
        return "\n".join(parts)
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    return path.read_text(encoding="utf-8", errors="replace")


def guess_basics(text: str) -> dict:
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    phone = re.search(r"\+?\d[\d \-/()]{7,}\d", text)
    linkedin = re.search(r"linkedin\.com/in/[\w\-]+", text, re.I)
    github = re.search(r"github\.com/[\w\-]+", text, re.I)
    first = next((ln.strip(" *#") for ln in text.splitlines() if ln.strip()), "")
    name = first if 1 <= len(first.split()) <= 5 and not re.search(r"\d|@", first) else ""
    skills = [t for t in TECH_TERMS if re.search(r"(?<![A-Za-z])" + re.escape(t) + r"(?![A-Za-z])", text, re.I)]
    return {"name": name.title() if name.isupper() else name, "email": email.group(0) if email else "",
            "phone": phone.group(0).strip() if phone else "",
            "linkedin": linkedin.group(0) if linkedin else "", "github": github.group(0) if github else "",
            "skills": skills}


# ---------------------------------------------------------------------------
def _parse_md(text: str) -> List[Tuple[str, str]]:
    blocks: List[Tuple[str, str]] = []
    para: List[str] = []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for line in text.splitlines():
        s = line.rstrip()
        if not s.strip():
            flush()
            continue
        m = re.match(r"^(#{1,3})\s+(.*)", s)
        if m:
            flush()
            blocks.append((f"h{len(m.group(1))}", m.group(2)))
        elif re.match(r"^\s*[-*•]\s+", s):
            flush()
            blocks.append(("li", re.sub(r"^\s*[-*•]\s+", "", s)))
        else:
            para.append(s.strip())
    flush()
    return blocks


def _runs(text: str) -> List[Tuple[str, bool, bool]]:
    out = []
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], True, False))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            out.append((part[1:-1], False, True))
        else:
            out.append((part, False, False))
    return out


def to_docx(md_path: Path, out: Path) -> Path:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for kind, text in _parse_md(md_path.read_text(encoding="utf-8")):
        if kind.startswith("h"):
            d.add_heading(text, level=int(kind[1]))
            continue
        p = d.add_paragraph(style="List Bullet" if kind == "li" else None)
        for t, b, i in _runs(text):
            r = p.add_run(t)
            r.bold, r.italic = b, i
    d.save(str(out))
    return out


def to_pdf(md_path: Path, out: Path) -> Path:
    from fpdf import FPDF

    def safe(s: str) -> str:
        s = s.replace("–", "-").replace("—", "-").replace("’", "'").replace("“", '"').replace("”", '"')
        return s.encode("latin-1", errors="replace").decode("latin-1")

    pdf = FPDF(format="A4")
    pdf.set_margins(20, 18, 20)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    sizes = {"h1": 16, "h2": 13, "h3": 11.5}
    for kind, text in _parse_md(md_path.read_text(encoding="utf-8")):
        if kind.startswith("h"):
            pdf.set_font("Helvetica", "B", sizes[kind])
            pdf.multi_cell(0, 7, safe(text.replace("**", "")))
            pdf.ln(1)
            continue
        pdf.set_font("Helvetica", "", 10.5)
        prefix = "- " if kind == "li" else ""
        # fpdf2 markdown uses __text__ for italics, markdown files use *text*
        text = re.sub(r"(?<![*\w])\*(?!\*)([^*\n]+?)\*(?!\*)", r"__\1__", text)
        md = safe(prefix + text)
        pdf.multi_cell(0, 5.5, md, markdown=True)
        pdf.ln(1.5 if kind == "li" else 3)
    pdf.output(str(out))
    return out
