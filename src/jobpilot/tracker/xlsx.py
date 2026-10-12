"""Local Excel tracker (works in Excel, LibreOffice, Numbers and Google Sheets import)."""
from __future__ import annotations

from typing import Any, List, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..config import Workspace
from ..models import APPLICATION_COLUMNS, JOB_COLUMNS, RUNLOG_COLUMNS, STATUSES, WATCHLIST_COLUMNS
from .base import DASH_START_ROW, HOW_TO, SHEETS, Tracker, dashboard_cells

NAVY = "1F3864"
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
BODY_FONT = Font(name="Arial", size=10)
GREEN = PatternFill("solid", fgColor="C8E6C9")
AMBER = PatternFill("solid", fgColor="FFF2CC")
ORANGE = PatternFill("solid", fgColor="F8CBAD")
RED = PatternFill("solid", fgColor="F4CCCC")
BLUE = PatternFill("solid", fgColor="CFE2F3")

WIDTHS = {"ID": 8, "Status": 12, "Fit": 8, "Score": 7, "Company": 24, "Title": 48, "Posting Ref": 18, "Type": 12,
          "Location": 18, "Work Mode": 11, "Posting Language": 10, "Language Required": 12, "Posted Date": 12,
          "Date Found": 12, "Source": 16, "Link": 14, "Notes": 40, "Applied Date": 12, "Last Update": 12,
          "Job ID": 8, "Channel": 14, "Documents Sent": 18, "Contact Person": 18, "Last Response Date": 13,
          "Next Step": 24, "Follow-up Due": 12, "Mail Reference": 30, "Interview Date": 14,
          "Run Time": 17, "Task": 14, "New Postings": 10, "Updated Records": 10, "Sources OK": 30,
          "Sources Failed": 30, "Active": 8, "ATS": 14, "Slug": 18, "Why": 40, "Last Checked": 12}


def _header(ws, cols: List[str]) -> None:
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c, 14)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


class XlsxTracker(Tracker):
    def __init__(self, ws: Workspace):
        super().__init__(ws)
        self.path = ws.tracker_path
        if not self.path.exists():
            self.path.parent.mkdir(parents=True, exist_ok=True)
            create_workbook(str(self.path), ws.get("person.name", ""))
        self.wb = load_workbook(self.path)
        for name in SHEETS:
            if name not in self.wb.sheetnames:
                raise RuntimeError(f"Tracker is missing the '{name}' tab. Rename it back or delete the file to recreate.")

    def _rows(self, sheet: str) -> List[List[Any]]:
        return [list(r) for r in self.wb[sheet].iter_rows(values_only=True)]

    def _style_row(self, sheet, r: int, ncols: int) -> None:
        for c in range(1, ncols + 1):
            sheet.cell(row=r, column=c).font = BODY_FONT

    def _append(self, sheet: str, rows: List[List[Any]]) -> None:
        sh = self.wb[sheet]
        # append after the last non empty row (data validation ranges make max_row unreliable)
        last = 1
        for r in range(sh.max_row, 0, -1):
            if any(sh.cell(row=r, column=c).value not in (None, "") for c in range(1, min(sh.max_column, 6) + 1)):
                last = r
                break
        for row in rows:
            last += 1
            for c, v in enumerate(row, 1):
                sh.cell(row=last, column=c, value=v)
            self._style_row(sh, last, len(row))
            if sheet == "Jobs":
                link_cell = sh.cell(row=last, column=JOB_COLUMNS.index("Link") + 1)
                if link_cell.value:
                    link_cell.hyperlink = str(link_cell.value)
                    link_cell.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    def _update_cells(self, sheet: str, updates: List[Tuple[int, int, Any]]) -> None:
        sh = self.wb[sheet]
        for r, c, v in updates:
            sh.cell(row=r, column=c + 1, value=v)

    def _replace_sheet(self, sheet: str, header: List[str], rows: List[List[Any]]) -> None:
        sh = self.wb[sheet]
        if sh.max_row > 1:
            sh.delete_rows(2, sh.max_row)
        _header(sh, header)
        for i, row in enumerate(rows, 2):
            for c, v in enumerate(row, 1):
                cell = sh.cell(row=i, column=c, value=v)
                cell.font = BODY_FONT
            if sheet == "Review Queue" and "Link" in header:
                lc = sh.cell(row=i, column=header.index("Link") + 1)
                if lc.value:
                    lc.hyperlink = str(lc.value)
                    lc.font = Font(name="Arial", size=10, color="0563C1", underline="single")

    def save(self) -> None:
        try:
            self.wb.save(self.path)
        except PermissionError as exc:
            raise RuntimeError(f"Could not save {self.path}. Close it in Excel and run again.") from exc


def create_workbook(path: str, owner: str = "") -> None:
    wb = Workbook()
    dash = wb.active
    dash.title = "Dashboard"
    for name in SHEETS[1:]:
        wb.create_sheet(name)

    # Dashboard
    dash["A1"] = "Job Search Dashboard"
    dash["A1"].font = Font(name="Arial", bold=True, size=16)
    dash["A2"] = f"{owner}. Built by JobPilot. Formulas update when you open the file." if owner else \
        "Built by JobPilot. Formulas update when you open the file."
    dash["A2"].font = Font(name="Arial", italic=True, size=10, color="595959")
    for i, (label, formula) in enumerate(dashboard_cells()):
        r = DASH_START_ROW + i
        a = dash.cell(row=r, column=1, value=label or None)
        if formula is None and label:
            a.font, a.fill = HEAD_FONT, HEAD_FILL
            dash.cell(row=r, column=2).fill = HEAD_FILL
        else:
            a.font = BODY_FONT
        if formula:
            b = dash.cell(row=r, column=2, value=formula)
            b.font = Font(name="Arial", bold=True, size=10)
            if "rate" in label.lower():
                b.number_format = "0.0%"
    for i, line in enumerate(HOW_TO):
        c = dash.cell(row=DASH_START_ROW + i, column=4, value=line)
        c.font = HEAD_FONT if i == 0 else BODY_FONT
        if i == 0:
            c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    dash.column_dimensions["A"].width = 28
    dash.column_dimensions["B"].width = 10
    dash.column_dimensions["C"].width = 3
    dash.column_dimensions["D"].width = 110

    _header(wb["Jobs"], JOB_COLUMNS)
    _header(wb["Applications"], APPLICATION_COLUMNS)
    _header(wb["Watchlist"], WATCHLIST_COLUMNS)
    _header(wb["Run Log"], RUNLOG_COLUMNS)
    _header(wb["Review Queue"], ["ID", "Fit", "Score", "Company", "Title", "Posting Ref", "Type", "Location",
                                 "Language Required", "Posted Date", "Link"])

    jobs = wb["Jobs"]
    n = 5000
    st = get_column_letter(JOB_COLUMNS.index("Status") + 1)
    fit = get_column_letter(JOB_COLUMNS.index("Fit") + 1)
    lang = get_column_letter(JOB_COLUMNS.index("Language Required") + 1)
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True, showErrorMessage=False)
    jobs.add_data_validation(dv)
    dv.add(f"{st}2:{st}{n}")
    dv2 = DataValidation(type="list", formula1='"High,Medium,Low"', allow_blank=True, showErrorMessage=False)
    jobs.add_data_validation(dv2)
    dv2.add(f"{fit}2:{fit}{n}")
    apps = wb["Applications"]
    ast = get_column_letter(APPLICATION_COLUMNS.index("Status") + 1)
    dv3 = DataValidation(type="list", formula1='"' + ",".join(STATUSES[2:]) + '"', allow_blank=True, showErrorMessage=False)
    apps.add_data_validation(dv3)
    dv3.add(f"{ast}2:{ast}1000")

    cf = jobs.conditional_formatting
    cf.add(f"{fit}2:{fit}{n}", CellIsRule(operator="equal", formula=['"High"'], fill=GREEN, font=Font(bold=True)))
    cf.add(f"{fit}2:{fit}{n}", CellIsRule(operator="equal", formula=['"Medium"'], fill=AMBER))
    cf.add(f"{fit}2:{fit}{n}", CellIsRule(operator="equal", formula=['"Low"'], fill=ORANGE))
    cf.add(f"{lang}2:{lang}{n}", CellIsRule(operator="equal", formula=['"Likely"'], fill=AMBER))
    cf.add(f"{lang}2:{lang}{n}", FormulaRule(formula=[f'OR(${lang}2="No",${lang}2="Optional")'], fill=GREEN))
    # Whole row colouring by status makes long lists easy to scan.
    last_col = get_column_letter(len(JOB_COLUMNS))
    cf.add(f"A2:{last_col}{n}", FormulaRule(formula=[f'OR(${st}2="Interview",${st}2="Offer")'], fill=GREEN))
    cf.add(f"A2:{last_col}{n}", FormulaRule(formula=[f'OR(${st}2="Applied",${st}2="Screening")'], fill=BLUE))
    cf.add(f"A2:{last_col}{n}", FormulaRule(formula=[f'${st}2="Rejected"'], fill=RED))
    cf.add(f"A2:{last_col}{n}", FormulaRule(formula=[f'OR(${st}2="Skipped",${st}2="Ghosted",${st}2="Withdrawn")'],
                                            font=Font(color="999999")))
    jobs.auto_filter.ref = f"A1:{last_col}1"
    jobs.freeze_panes = "F2"
    wb.save(path)
