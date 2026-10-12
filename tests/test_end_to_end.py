from datetime import datetime

import pytest
import yaml
from openpyxl import load_workbook

from jobpilot import cli, config
from jobpilot.inbox import Mail, classify, run_inbox
from jobpilot.search import run_search
from jobpilot.tracker import open_tracker

from conftest import FakeHttp, load

ROUTES = {"arbeitnow.com": load("arbeitnow.json"), "arbeitsagentur.de": load("bundesagentur.json"),
          "greenhouse.io": load("greenhouse.json")}
NOW = datetime(2026, 10, 12, 6, 47)


def test_search_writes_tracker_and_dedupes(workspace):
    res = run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    assert not res["failed"], res["failed"]
    ids = [r[0] for r in res["rows"]]
    assert ids[0] == "J0001" and len(ids) == len(set(ids)) >= 3
    t = open_tracker(workspace)
    assert len(t.jobs()) == len(ids)
    # review queue rebuilt, High first
    wb = load_workbook(workspace.tracker_path)
    q = [r for r in wb["Review Queue"].iter_rows(min_row=2, values_only=True) if r[0]]
    assert len(q) == len(ids) and q[0][1] == "High"
    assert wb["Watchlist"]["A2"].value == "Acme AI"
    assert wb["Run Log"]["B2"].value == "Job search"
    # second run: nothing new
    res2 = run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    assert res2["rows"] == [] and res2["stats"]["duplicate"] >= len(ids)


def test_save_is_atomic_and_keeps_backup(workspace):
    run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    t = open_tracker(workspace)
    t.update_job("J0001", {"Status": "Shortlisted"}, "2026-10-12")
    t.save()
    bak = workspace.tracker_path.with_name(workspace.tracker_path.stem + ".bak.xlsx")
    assert bak.exists(), "last good version kept as .bak"
    assert not workspace.tracker_path.with_name(workspace.tracker_path.stem + ".tmp.xlsx").exists()
    # the backup is the version from before this save
    old = load_workbook(bak)["Jobs"]
    new = load_workbook(workspace.tracker_path)["Jobs"]
    assert new["B2"].value == "Shortlisted" and old["B2"].value == "Found"


def test_failing_source_does_not_stop_run(workspace):
    routes = dict(ROUTES)
    routes["arbeitsagentur.de"] = RuntimeError("boom")
    res = run_search(workspace, http=FakeHttp(routes), now=NOW, quiet=True)
    assert any("bundesagentur" in f for f in res["failed"]) and res["rows"]


def test_cli_track_add_show_status(workspace, capsys):
    run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    w = ["--workspace", str(workspace.path)]
    assert cli.main(w + ["track", "J0001", "applied", "--note", "LinkedIn, CV and cover letter"]) == 0
    t = open_tracker(config.Workspace(workspace.path))
    j = t.find_job("J0001")
    assert j["Status"] == "Applied" and j["Applied Date"]
    a = t.find_application("J0001")
    assert a and a["Status"] == "Applied" and a["Follow-up Due"]
    # no downgrade without --force
    cli.main(w + ["track", "J0001", "Found"])
    assert open_tracker(config.Workspace(workspace.path)).find_job("J0001")["Status"] == "Applied"
    assert cli.main(w + ["add", "--company", "Self Found GmbH", "--title", "Working Student Backend (m/w/d)",
                         "--link", "https://example.com/job/1", "--description", "Python APIs, English team"]) == 0
    assert cli.main(w + ["show", "J0001"]) == 0
    assert cli.main(w + ["status"]) == 0
    out = capsys.readouterr().out
    assert "Self Found GmbH" in out or "Added" in out
    assert cli.main(w + ["track", "J9999", "Applied"]) == 1


def test_inbox_updates_statuses(workspace):
    run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    t = open_tracker(workspace)
    jobs = {j["Company"]: j for j in t.jobs()}
    glind = jobs["Glindway GmbH"]["ID"]
    acme = jobs["Acme AI"]["ID"]
    t.update_job(glind, {"Status": "Applied"}, "2026-10-01")
    t.upsert_application(glind, {"Applied Date": "2026-10-01", "Status": "Applied"})
    t.save()
    mails = [
        Mail("Glindway Recruiting <jobs@glindway.com>", "Your application: invitation to an interview",
             "Hi, we would like to invite you to an interview next week. Please pick a slot.", NOW.date(), "<1>"),
        Mail("Acme AI <no-reply@greenhouse.io>", "Thank you for applying to Acme AI",
             "We have received your application for Working Student Machine Learning. Ref REQ-2026-77", NOW.date(), "<2>"),
        Mail("StepStone <noreply@stepstone.de>", "Neue Jobs für dich", "Interview tips and 20 jobs", NOW.date(), "<3>"),
        Mail("Random <hello@random.org>", "Newsletter", "Nothing relevant", NOW.date(), "<4>"),
    ]
    res = run_inbox(workspace, mails=mails, today="2026-10-12", quiet=True)
    changes = {c[1]: (c[2], c[3]) for c in res["changes"]}
    assert changes["Glindway GmbH"] == ("Applied", "Interview")
    assert changes["Acme AI"] == ("Found", "Applied")
    assert len(changes) == 2
    t = open_tracker(workspace)
    assert t.find_application(acme)["Status"] == "Applied"
    assert t.find_job(glind)["Notes"].startswith("2026-10-12: interview mail")


def test_classifier_languages():
    m = lambda s, b: Mail("x@y.de", s, b, NOW.date(), "")  # noqa: E731
    assert classify(m("Ihre Bewerbung", "leider müssen wir Ihnen mitteilen, dass wir uns anders entschieden haben")) == "Rejected"
    assert classify(m("Einladung zum Vorstellungsgespräch", "Wir freuen uns")) == "Interview"
    assert classify(m("Coding challenge", "Please complete the HackerRank test")) == "Screening"
    assert classify(m("Vielen Dank für Ihre Bewerbung", "Eingang Ihrer Bewerbung bestätigt")) == "Applied"
    assert classify(m("Hello", "Lunch tomorrow?")) is None


def test_ghosting_rule(workspace):
    run_search(workspace, http=FakeHttp(ROUTES), now=NOW, quiet=True)
    t = open_tracker(workspace)
    t.update_job("J0001", {"Status": "Applied"}, "2026-09-01")
    t.upsert_application("J0001", {"Applied Date": "2026-09-15", "Status": "Applied"})
    t.update_job("J0002", {"Status": "Applied"}, "2026-08-01")
    t.upsert_application("J0002", {"Applied Date": "2026-08-01", "Status": "Applied"})
    t.save()
    res = run_inbox(workspace, mails=[], today="2026-10-12", quiet=True)
    t = open_tracker(workspace)
    assert t.find_application("J0001")["Next Step"] == "Send polite follow-up"
    assert t.find_job("J0002")["Status"] == "Ghosted"
    assert len(res["follow_ups"]) == 1
